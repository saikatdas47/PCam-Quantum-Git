from pathlib import Path
import json,re,hashlib
from docx import Document
from xml.etree import ElementTree as E
P=Path('paper/plos_one_word');D=Document(P/'PCam_ResNet50_Quantum_Manuscript.docx');text='\n'.join(p.text for p in D.paragraphs)
expected='A comparison of simulated quantum and classical classification heads on frozen ResNet50 features for PatchCamelyon'
assert D.paragraphs[0].text==D.core_properties.title==expected
assert 'controlled' not in expected.lower()
assert D.core_properties.author=='Saikat Das'
for value in ['Saikat Das','Department of Computer Science and Engineering','Ahsanullah University of Science and Technology, Dhaka, Bangladesh','saikatdasmain47@gmail.com','The author received no specific funding','The author declares no competing interests','Ethics approval and consent to participate: Not applicable.']:assert value in text,value
assert 'Acknowledgments' not in text
assert 'Authors and affiliations to be completed' not in text
links=[r.target_ref for r in D.part.rels.values() if r.reltype.endswith('/hyperlink')]
for target in ['mailto:saikatdasmain47@gmail.com','https://github.com/saikatdas47/PCam-Quantum-Git','https://github.com/basveeling/pcam#download','https://zenodo.org/records/2546921']:assert target in links,target
assert (P/'MANUSCRIPT.md').read_text().splitlines()[0]=='# '+expected
before=json.loads((P/'consistency_audit/before.json').read_text())
after=[[[c.text for c in row.cells] for row in t.rows] for t in D.tables]
assert before['tables']==after
assert len(D.tables)==11 and len(D.inline_shapes)==5
cap=[p.text for p in D.paragraphs if re.match(r'^Fig [1-5]\.',p.text)]
assert len(cap)==5 and [int(re.match(r'^Fig (\d)',s)[1]) for s in cap]==[1,2,3,4,5]
for n in range(1,6):
 # All figures have a prose callout as well as a caption.
 assert any(not p.text.startswith('Fig '+str(n)+'.') and (re.search(r'\bFig '+str(n)+r'\b',p.text) or (n in [2,3] and 'Figs 2 and 3' in p.text)) for p in D.paragraphs),n
assert 'methodology figure remains disabled' not in text.lower()
refs=re.findall(r'^(\d+)\. .+',text,re.M);assert list(map(int,refs))==list(range(1,26))
body=text[:text.index('\nReferences\n')];cited=set()
for group in re.findall(r'\[(\d+(?:,\d+)*)\]',body):cited.update(map(int,group.split(',')))
assert cited==set(range(1,26))
assert 'Brier score and log loss evaluated predicted probabilities.' in text
assert 'with epoch indexed from zero' in text
assert 'All neural head fits used CPU float64 with two PyTorch threads.' in text
# Figure IDs and labels are genuine editable graph elements, not a flattened bitmap.
x=E.parse('paper/PLOS_ONE_submission/figures/Fig1_methodology.drawio');vals='\n'.join(c.get('value','') for c in x.findall('.//mxCell'))
for val in ['Frozen ResNet50','16,384','262,144','PCA: 2,048 → 8','321 parameters','25 parameters','rate 0.003','100 epochs','101, 202, 303','17,090','15,678','129 whole slide','20,000','Exact simulation']:assert val in vals,val
assert len(body.split('Abstract\n',1)[1].split('\nIntroduction',1)[0].split())<=300
placeholders=[p.text for p in D.paragraphs if ('[Author' in p.text or '[Corresponding' in p.text or '[PUBLIC REPOSITORY' in p.text or '[Acknowledgments' in p.text)]
res={'title':expected,'title_and_metadata_match':True,'all_11_tables_unchanged':True,'figures':5,'figure_caption_order':[1,2,3,4,5],'all_figures_cited_in_prose':True,'references':25,'all_references_cited':True,'abstract_words':len(body.split('Abstract\n',1)[1].split('\nIntroduction',1)[0].split()),'diagram_scope_labels_match_text':True,'probability_and_threshold_metrics_separated':True,'explicit_epoch_index_and_head_runtime':True,'remaining_author_placeholders':placeholders,'sha256':hashlib.sha256((P/'PCam_ResNet50_Quantum_Manuscript.docx').read_bytes()).hexdigest()}
(P/'consistency_audit/document_checks.json').write_text(json.dumps(res,indent=2));print(json.dumps({k:v for k,v in res.items() if k!='remaining_author_placeholders'},indent=2))
